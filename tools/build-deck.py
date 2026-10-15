"""Build the editable eight-slide development handoff deck with python-pptx."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
root=Path(__file__).resolve().parents[1]
r=json.loads((root/'results/alwaz/access-eval.json').read_text())
c=json.loads((root/'results/alwaz/classifier-spike.json').read_text())['results']
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
BG='111B2B';FG='F1F5FA';MUTED='B6C3D7';GREEN='64DEB1'
def text(s,txt,x,y,w,h,size=24,color=FG,bold=False):
 box=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=box.text_frame;tf.word_wrap=True
 for i,line in enumerate(txt.split('\n')):
  para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.text=line;para.font.name='Aptos';para.font.size=Pt(size);para.font.bold=bold;para.font.color.rgb=RGBColor.from_string(color);para.space_after=Pt(14)
 return box
def slide(title,sub=None):
 s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 text(s,title,.65,.48,12,1,36,bold=True)
 if sub:text(s,sub,.68,1.5,12,.7,19,MUTED)
 text(s,str(len(p.slides)),12.1,6.95,.5,.3,12,MUTED)
 return s
def table(s,headers,rows,y=2):
 sh=s.shapes.add_table(len(rows)+1,len(headers),Inches(.7),Inches(y),Inches(11.9),Inches(.65*(len(rows)+1))).table
 for ri,row in enumerate([headers]+rows):
  for ci,val in enumerate(row):
   cell=sh.cell(ri,ci);cell.text=str(val);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('243449' if ri==0 else BG)
   for para in cell.text_frame.paragraphs:
    para.font.name='Aptos';para.font.size=Pt(19);para.font.bold=ri==0;para.font.color.rgb=RGBColor.from_string(GREEN if ri==0 else FG)
 return sh
s=slide('LeakShield', 'Access and approvals · Alwaz · CyberCUP 6.0')
text(s,'Control which documents\nreach the model.',.7,2.65,11.9,2,44,bold=True)
text(s,'Development evidence · 15 October 2026',.7,5.9,12,.6,21,MUTED)
s.notes_slide.notes_text_frame.text='This deck reports Alwaz’s implemented scope. The full gateway and final frozen benchmark are team integration dependencies.'
s=slide('Architecture','ACLs apply inside MongoDB and again before provider assembly.')
# Native editable architecture is an explicit task deliverable.
labels=['Authenticated\nprincipal','Mongo query\nwith ACL','Fresh ACL +\ncontent hash','Outbound scan\ncallback','Provider +\noutput checks']
for i,label in enumerate(labels):
 x=.7+i*2.48
 shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.7),Inches(2.12),Inches(1.3));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('243449');shape.line.color.rgb=RGBColor.from_string(GREEN)
 text(s,label,x+.08,2.88,1.96,1,19)
 if i<4:
  con=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x+2.12),Inches(3.35),Inches(x+2.48),Inches(3.35));con.line.color.rgb=RGBColor.from_string(GREEN)
text(s,'Approval → current policy and ACL check → outbound scan again',.75,4.8,11.8,.9,24,GREEN)
text(s,'The provider and final scanner callbacks are supplied by the team gateway.',.75,6.0,11.8,.6,19,MUTED)
s=slide('Identity and provider permission','Read access and external transmission are separate decisions.')
table(s,['Scenario','Expected boundary'],[['Asha requests salary bands','Deny; no restricted title'],['Ravi requests salary bands','Permit HR retrieval'],['Admin requests board minutes','Deny; no document bypass'],['CEO requests board minutes','Read locally; deny cloud send'],['Meera requests Garuda','Require admin approval']])
s=slide('Quarantine and approvals')
text(s,'Ingest source → scan hidden carriers → quarantine whole document',.75,1.9,11.8,1.1,27,GREEN)
text(s,'Comments, hidden CSS, white text, zero-width characters and alt/title attributes are checked before rendering.',.75,3.1,11.6,1.1,24)
text(s,'Pending requests use AES-GCM in memory for ten minutes. Mongo stores metadata and request hash. Approval is single use and rechecks current permissions.',.75,4.7,11.6,1.5,24)
s=slide('Measured access evidence','Real local MongoDB · simulated provider · development fixtures')
table(s,['Measurement','Observed'],[['Vitest security suite','20 passed · 0 failed'],['Expected access actions',f"{r['passed']} / {r['cases']}"],['Synthetic corpus','10 documents · 2 quarantined'],['Local access p50 / p95',f"{r['p50_ms']:.3f} / {r['p95_ms']:.3f} ms"]])
text(s,'Excludes real detectors, model/network time and human approval wait. This is not the final frozen gateway evaluation.',.75,5.9,11.7,.9,19,MUTED)
s.notes_slide.notes_text_frame.text='Source: results/alwaz/access-eval.json. Access-only fixtures are not detector precision/recall or full-team leak metrics.'
s=slide('Classifier feasibility','Pinned ProtectAI DeBERTa model · local Node ONNX · English only')
table(s,['Variant','Warm p50','Warm p95'],[[x['dtype'],f"{x['p50_ms']:.2f} ms",f"{x['p95_ms']:.2f} ms"] for x in c],2.2)
text(s,'Node works. Full precision loads directly upstream; int8 was generated locally. A sidecar was unnecessary.',.75,4.5,11.7,1,25,GREEN)
text(s,'Six short development prompts, three rounds; first round excluded. No long-window or Hinglish accuracy claim.',.75,5.95,11.7,.9,19,MUTED)
s.notes_slide.notes_text_frame.text='Model and Apache-2.0 license: https://huggingface.co/protectai/deberta-v3-base-prompt-injection-v2 . Revision 90c9989b1a342275dd0d1a95aad283c04e075671. Measurements: results/alwaz/classifier-spike.json. Shared-process RSS is unsuitable for isolated memory comparisons.'
s=slide('DPDP Act: control mapping','Technical alignment; this demo does not establish legal compliance.')
table(s,['Control','Act reference'],[['Limit outbound personal data','Section 6(1): consent and necessary data'],['ACLs and disclosure blocking','Section 8(5): security safeguards'],['Incident evidence supports review','Section 8(6): breach intimation'],['Expiry and deletion controls','Section 8(7): erasure obligations']])
text(s,'These provisions have phased commencement; as of 15 Oct 2026, the cited substantive duties are scheduled for 13 May 2027.',.75,5.9,11.7,.9,19,MUTED)
s.notes_slide.notes_text_frame.text='Primary source: https://www.indiacode.nic.in/bitstream/123456789/22037/1/a2023-22.pdf ; sections 6(1), 8(5)-(7), commencement footnote at page 5. Controls assist implementation but do not satisfy notice, consent, notification or erasure workflows on their own.'
s=slide('Integration limits and next gates')
text(s,'Ready: access modules, approval backend, seed, fixtures, Node classifier spike and reproducible tests.',.75,1.9,11.8,1.2,27,GREEN)
text(s,'Team integration: real gateway scans and vault, UI wiring, audit privacy, Groq/Ollama baseline and frozen evaluation.',.75,3.5,11.8,1.3,25)
text(s,'Approved content reaches Groq. Fingerprints catch copying, not paraphrases. Team Groq settings still need account verification.',.75,5.25,11.8,1.1,23,MUTED)
s.notes_slide.notes_text_frame.text='Groq policy source: https://console.groq.com/docs/your-data . Team-specific ZDR and rate limits have not been verified. Implementation published for review: https://github.com/TeamElara/Aevia-AI/pull/1 .'
p.save(root/'docs/Alwaz-LeakShield.pptx')
print('Created 8 editable slides')
